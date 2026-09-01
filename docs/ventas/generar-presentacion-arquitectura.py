# -*- coding: utf-8 -*-
"""Presentación técnica de arquitectura MediCore (para arquitectos de software)."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("MediCore-Arquitectura-Tecnica.pptx")

NAVY = RGBColor(0x0B, 0x2C, 0x3D)
TEAL = RGBColor(0x0D, 0x7A, 0x7A)
TEAL_LIGHT = RGBColor(0xE6, 0xF4, 0xF4)
CORAL = RGBColor(0xC4, 0x5C, 0x26)
AMBER = RGBColor(0xB4, 0x7E, 0x0A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1A, 0x2A, 0x32)
MUTED = RGBColor(0x5A, 0x6B, 0x73)
LINE = RGBColor(0xD0, 0xDC, 0xE0)
CARD_BG = RGBColor(0xF7, 0xFA, 0xFB)
CODE_BG = RGBColor(0x12, 0x2A, 0x38)


def set_run(run, text, size=16, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_textbox(slide, left, top, width, height, text, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    set_run(r, text, size, bold, color, font)
    return box


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def banner(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.05))
    fill_shape(bar, NAVY)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.05), Inches(13.333), Inches(0.07))
    fill_shape(accent, TEAL)
    add_textbox(slide, Inches(0.5), Inches(0.22), Inches(12.2), Inches(0.45), title, 24, True, WHITE)
    if subtitle:
        add_textbox(slide, Inches(0.5), Inches(0.62), Inches(12.2), Inches(0.32), subtitle, 12, False, RGBColor(0xB8, 0xD4, 0xD8))


def footer(slide, page, total):
    add_textbox(slide, Inches(0.5), Inches(7.15), Inches(10), Inches(0.28), "MediCore · Arquitectura técnica · docs/analisis/12 + 03", 10, False, MUTED)
    add_textbox(slide, Inches(11.2), Inches(7.15), Inches(1.5), Inches(0.28), f"{page} / {total}", 10, False, MUTED, PP_ALIGN.RIGHT)


def card(slide, left, top, width, height, title, body, title_color=TEAL, body_size=12):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill_shape(shp, CARD_BG)
    shp.line.color.rgb = LINE
    shp.adjustments[0] = 0.06
    add_textbox(slide, left + Inches(0.15), top + Inches(0.12), width - Inches(0.3), Inches(0.35), title, 13, True, title_color)
    add_textbox(slide, left + Inches(0.15), top + Inches(0.48), width - Inches(0.3), height - Inches(0.6), body, body_size, False, INK)


def bullets(slide, left, top, width, height, items, size=14):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(6)
        r = p.add_run()
        set_run(r, f"•  {item}", size, False, INK)


def node_box(slide, left, top, width, height, title, body, fill=TEAL, title_size=13):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill_shape(shp, fill)
    shp.adjustments[0] = 0.08
    add_textbox(slide, left + Inches(0.1), top + Inches(0.12), width - Inches(0.2), Inches(0.35), title, title_size, True, WHITE, PP_ALIGN.CENTER)
    if body:
        add_textbox(slide, left + Inches(0.1), top + Inches(0.45), width - Inches(0.2), height - Inches(0.55), body, 11, False, RGBColor(0xD8, 0xEE, 0xEE), PP_ALIGN.CENTER)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slides = []

    def new():
        sl = prs.slides.add_slide(blank)
        slides.append(sl)
        return sl

    # ---- 1 Portada ----
    s = new()
    fill_shape(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height), NAVY)
    fill_shape(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(5.7), prs.slide_width, Inches(1.8)), TEAL)
    add_textbox(s, Inches(0.8), Inches(1.5), Inches(11.5), Inches(0.6), "MediCore", 44, True, WHITE, font="Consolas")
    add_textbox(s, Inches(0.8), Inches(2.3), Inches(11.5), Inches(0.5), "Propuesta de arquitectura", 28, True, RGBColor(0xA8, 0xD8, 0xD8))
    add_textbox(
        s,
        Inches(0.8),
        Inches(3.1),
        Inches(11.5),
        Inches(1.4),
        "Core central + SPA/PWA · Offline por cola de dispositivo\nSin Edge · Outbox · Multi-tenant SP · DGIS/SINBA fijo",
        16,
        False,
        RGBColor(0xB0, 0xC8, 0xD0),
    )
    add_textbox(
        s,
        Inches(0.8),
        Inches(5.95),
        Inches(11.5),
        Inches(1.1),
        "Audiencia: arquitectos de software / tech leads\nFuente de verdad: docs/analisis/12-propuesta-final.md  ·  docs/analisis/03-arquitectura-propuesta.md\nEstado: 2026-08-23",
        13,
        False,
        WHITE,
    )

    # ---- 2 Agenda ----
    s = new()
    banner(s, "Agenda técnica", "Qué se decide y qué se descarta")
    items = [
        ("01", "Premisa clínica → arquitectura"),
        ("02", "Topología vigente (sin Edge)"),
        ("03", "Stack, capas y bounded contexts"),
        ("04", "Asimetría escritura / lectura"),
        ("05", "Offline, caché y dispositivos"),
        ("06", "Outbox, integraciones y BI"),
        ("07", "Multi-tenant, SPs, seguridad"),
        ("08", "Hospedaje, on-prem y ADRs"),
        ("09", "Contrato de no-promesas"),
    ]
    for i, (n, t) in enumerate(items):
        col, row = i % 3, i // 3
        left = Inches(0.5) + Inches(col * 4.2)
        top = Inches(1.5) + Inches(row * 1.7)
        card(s, left, top, Inches(4.0), Inches(1.45), n, t, body_size=14)

    # ---- 3 Premisa ----
    s = new()
    banner(s, "Premisa que define la arquitectura", "No parte de un requisito técnico")
    q = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.35), Inches(12.3), Inches(2.0))
    fill_shape(q, TEAL_LIGHT)
    q.line.color.rgb = TEAL
    add_textbox(
        s,
        Inches(0.75),
        Inches(1.55),
        Inches(11.8),
        Inches(1.6),
        "“Un paciente en urgencias vive… el modo offline existe para que no haya\nun ‘¿está caído el sistema?’ con alguien en fila. Incluso si no sabemos su nombre.”",
        16,
        True,
        NAVY,
        PP_ALIGN.CENTER,
    )
    for i, (t, b) in enumerate(
        [
            ("Criterio rector", "Ante conflicto atención vs. administrativo → gana la atención. Reconciliación después, con rastro auditable."),
            ("Límite innegociable", "No se relaja integridad clínica: append-only, occurred_at / recorded_at, autoría verificable."),
            ("Consecuencia técnica", "La resiliencia vive en el dispositivo (cola durable), no en un servidor Edge en la clínica."),
        ]
    ):
        card(s, Inches(0.5) + Inches(i * 4.2), Inches(3.65), Inches(4.0), Inches(2.85), t, b, body_size=13)

    # ---- 4 Topología ----
    s = new()
    banner(s, "Topología vigente", "ADR-001 reformulado · Core central + SPA/PWA")
    # clients
    node_box(s, Inches(0.5), Inches(1.8), Inches(3.2), Inches(2.2), "Estaciones", "SPA / PWA (mismo artefacto)\nIndexedDB: cola + caché\nSignalR cuando hay red", TEAL)
    # arrow
    add_textbox(s, Inches(3.7), Inches(2.5), Inches(0.8), Inches(0.5), "HTTPS", 12, True, MUTED, PP_ALIGN.CENTER)
    node_box(s, Inches(4.5), Inches(1.5), Inches(4.0), Inches(2.8), "Core API (.NET)", "Monolito modular\nClean Architecture\nJWT + TenantId\nApiResponse<T>", NAVY)
    add_textbox(s, Inches(8.5), Inches(2.5), Inches(0.7), Inches(0.5), "SPs", 12, True, MUTED, PP_ALIGN.CENTER)
    node_box(s, Inches(9.2), Inches(1.8), Inches(3.5), Inches(2.2), "SQL Server", "1 BD multi-tenant\nPersistencia solo SPs\nsp_{Entity}_{Action}", TEAL)
    # worker
    node_box(s, Inches(4.5), Inches(4.6), Inches(4.0), Inches(1.9), "Worker (outbox)", "DGIS/SINBA (siempre)\nCFDI / FHIR / RENAPO (flags)\nNo tumba la consulta", CORAL)
    add_textbox(s, Inches(0.5), Inches(4.7), Inches(3.8), Inches(1.6), "Prohibido en MVP\n• Edge por sucursal\n• Sync expediente Edge↔Core\n• Broker (Kafka/Rabbit)\n• Microservicios / BD-por-tenant / Redis", 12, False, MUTED)

    # ---- 5 Descarta Edge ----
    s = new()
    banner(s, "Qué se descartó (y por qué importa)", "Docs 05–07 pueden tener texto histórico; prevalece doc 12")
    for i, (t, b, c) in enumerate(
        [
            ("Edge por clínica", "Réplica + sync bidireccional del expediente. Conflicto con doc 2 (una BD central) y complejidad de conflicto clínico.", CORAL),
            ("App BI separada", "Sustituida por módulo features/bi en la app principal (solo agregados, lazy, auditoría).", AMBER),
            ("DGIS como feature flag", "Capacidad fija del producto (ADR-023). Sin red → outbox; el módulo no se apaga.", TEAL),
            ("Operación idéntica offline", "No se promete. Continuidad por estación + degradación declarada de coordinación compartida.", NAVY),
        ]
    ):
        col, row = i % 2, i // 2
        card(s, Inches(0.5) + Inches(col * 6.35), Inches(1.4) + Inches(row * 2.6), Inches(6.1), Inches(2.4), t, b, c, 13)

    # ---- 6 Stack ----
    s = new()
    banner(s, "Stack y capas", "Alineado a Clean Architecture del doc 2")
    layers = [
        ("Api", "Controllers, auth, middleware, SignalR hubs"),
        ("Business", "Casos de uso, políticas, bounded contexts"),
        ("DataAccess", "Solo invocación de SPs; sin SQL ad-hoc de negocio"),
        ("Models / Common", "DTOs, ApiResponse, cross-cutting"),
        ("Worker", "Proceso separado: drenado de outbox externo"),
        ("Frontend", "React 19 + TS + Vite + Tailwind · PWA"),
    ]
    for i, (t, b) in enumerate(layers):
        col, row = i % 3, i // 3
        card(s, Inches(0.45) + Inches(col * 4.25), Inches(1.4) + Inches(row * 2.55), Inches(4.05), Inches(2.35), t, b, body_size=14)

    # ---- 7 Write/Read ----
    s = new()
    banner(s, "Asimetría escritura / lectura", "Núcleo del diseño · ADR-014")
    # Write column
    w = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(1.35), Inches(6.1), Inches(5.4))
    fill_shape(w, CARD_BG)
    w.line.color.rgb = TEAL
    add_textbox(s, Inches(0.7), Inches(1.5), Inches(5.6), Inches(0.4), "ESCRIBIR — siempre local", 16, True, TEAL)
    bullets(
        s,
        Inches(0.7),
        Inches(2.05),
        Inches(5.6),
        Inches(4.4),
        [
            "IdempotencyKey al capturar (no al enviar)",
            "Persistir primero en cola IndexedDB",
            "UI confirma contra local — nunca contra la red",
            "Drenador en background → API",
            "Reintentos con backoff; falla permanente no bloquea cola",
            "Prohibido: if (online) remote else local por módulo",
        ],
        14,
    )
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.35), Inches(6.1), Inches(5.4))
    fill_shape(r, CARD_BG)
    r.line.color.rgb = NAVY
    add_textbox(s, Inches(7.05), Inches(1.5), Inches(5.6), Inches(0.4), "LEER — servidor con enlace", 16, True, NAVY)
    bullets(
        s,
        Inches(7.05),
        Inches(2.05),
        Inches(5.6),
        Inches(4.4),
        [
            "Online: servidor = verdad",
            "Offline: última foto + antigüedad visible (SC-09)",
            "Live (SignalR): cola urgencias / sala / turnos",
            "Caché mínima; no padrón ni HC completa",
            "Logout: purga lecturas; NO borra cola de salida",
        ],
        14,
    )

    # ---- 8 Cache levels ----
    s = new()
    banner(s, "Cuatro niveles de caché + dispositivos", "ADR-017: solo equipos registrados cachean clínico")
    levels = [
        ("N1", "Shell / catálogos / config / marca / permisos", "Completo, siempre", TEAL),
        ("N2", "Conjunto clínico de seguridad (día, cola, recientes)", "Mínimo + caducidad", TEAL),
        ("N3", "Estado compartido (sala, existencias, turnos)", "Última foto — nunca verdad", AMBER),
        ("N4", "Resto", "Por demanda", MUTED),
    ]
    for i, (n, t, p, c) in enumerate(levels):
        top = Inches(1.35) + Inches(i * 1.05)
        nb = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), top, Inches(1.0), Inches(0.85))
        fill_shape(nb, c)
        add_textbox(s, Inches(0.5), top + Inches(0.22), Inches(1.0), Inches(0.45), n, 16, True, WHITE, PP_ALIGN.CENTER)
        add_textbox(s, Inches(1.7), top + Inches(0.1), Inches(7.5), Inches(0.35), t, 14, True, INK)
        add_textbox(s, Inches(1.7), top + Inches(0.45), Inches(7.5), Inches(0.35), p, 13, False, MUTED)
    note = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.4), Inches(1.35), Inches(3.4), Inches(5.15))
    fill_shape(note, NAVY)
    add_textbox(s, Inches(9.6), Inches(1.6), Inches(3.0), Inches(0.4), "Dispositivos", 16, True, WHITE)
    add_textbox(
        s,
        Inches(9.6),
        Inches(2.2),
        Inches(3.0),
        Inches(3.8),
        "Registrados por admin:\ncola + caché clínica.\n\nNavegador no registrado:\nsolo online.\n\nLogout:\npurga lecturas;\nconserva outbox local.",
        13,
        False,
        RGBColor(0xC8, 0xE0, 0xE4),
    )

    # ---- 9 Failure domains ----
    s = new()
    banner(s, "Dominios de falla (contrato)", "Qué cubre el diseño y qué no")
    fails = [
        ("F1", "Internet caído", "Cola local + caché + antigüedad"),
        ("F2", "API/BD caída o deploy", "Igual F1; deploy escalonado solo en dedicado"),
        ("F3", "Estación / navegador", "IndexedDB durable"),
        ("F4", "Corte eléctrico", "Papel + captura diferida (occurred≠recorded)"),
        ("F5", "Pérdida DC", "Backups/DR proveedor — no hay Edge que absorba"),
        ("F6", "Corrupción lógica", "Append-only + PITR + cadena integridad"),
        ("F7", "Compromiso / ransomware", "MFA, segmentación, backups inmutables, devices"),
    ]
    for i, (n, t, b) in enumerate(fails):
        col, row = i % 4, i // 4
        if i == 6:
            left = Inches(0.5) + Inches(1.5)
        else:
            left = Inches(0.4) + Inches(col * 3.2)
        top = Inches(1.35) + Inches(row * 2.7)
        card(s, left, top, Inches(3.05), Inches(2.45), f"{n}  {t}", b, body_size=13)

    # ---- 10 Offline by domain ----
    s = new()
    banner(s, "Offline por dominio", "Matriz vigente (doc 03 §4)")
    rows = [
        ("Triage / SV", "Sí", "Append-only; no bloquear crítico"),
        ("Nota clínica", "Sí", "Firma local + sello sync; inmutable"),
        ("Urgencias / no ID", "Sí", "Identidad progresiva (doc 08)"),
        ("Receta", "Condicionado", "Catálogo + alergias en caché"),
        ("Agenda", "Lectura", "Sin árbitro local"),
        ("Farmacia", "Sí*", "Existencia negativa; *estupefacientes pendiente"),
        ("Caja", "Sí", "Recibo provisional; CFDI diferido"),
        ("DGIS/SINBA", "Genera+outbox", "Módulo nunca off"),
        ("CFDI/FHIR/RENAPO", "Flags", "Sin red o flag off → no bloquean"),
    ]
    # header
    add_textbox(s, Inches(0.5), Inches(1.3), Inches(3.5), Inches(0.35), "Dominio", 12, True, MUTED)
    add_textbox(s, Inches(4.1), Inches(1.3), Inches(2.2), Inches(0.35), "Sin backend", 12, True, MUTED)
    add_textbox(s, Inches(6.5), Inches(1.3), Inches(6.2), Inches(0.35), "Notas", 12, True, MUTED)
    for i, (d, o, n) in enumerate(rows):
        top = Inches(1.65) + Inches(i * 0.55)
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), top, Inches(12.4), Inches(0.5))
        fill_shape(bg, CARD_BG if i % 2 == 0 else WHITE)
        bg.line.color.rgb = LINE
        add_textbox(s, Inches(0.55), top + Inches(0.08), Inches(3.5), Inches(0.35), d, 12, True, INK)
        add_textbox(s, Inches(4.1), top + Inches(0.08), Inches(2.2), Inches(0.35), o, 12, False, TEAL)
        add_textbox(s, Inches(6.5), top + Inches(0.08), Inches(6.2), Inches(0.35), n, 12, False, MUTED)

    # ---- 11 Outbox ----
    s = new()
    banner(s, "EDA acotada + Outbox", "ADR-018 · Simetría cliente↔Core↔exterior")
    node_box(s, Inches(0.5), Inches(2.0), Inches(3.5), Inches(2.4), "Cola cliente", "IndexedDB → API\nIdempotencyKey\nReintentos / rechazo diferido", TEAL)
    add_textbox(s, Inches(4.1), Inches(2.9), Inches(0.7), Inches(0.4), "→", 22, True, MUTED, PP_ALIGN.CENTER)
    node_box(s, Inches(4.8), Inches(2.0), Inches(3.7), Inches(2.4), "Core + Outbox TX", "Misma TX que el SP\nHechos append-only\nSin broker en MVP", NAVY)
    add_textbox(s, Inches(8.6), Inches(2.9), Inches(0.7), Inches(0.4), "→", 22, True, MUTED, PP_ALIGN.CENTER)
    node_box(s, Inches(9.3), Inches(2.0), Inches(3.5), Inches(2.4), "Mundo exterior", "DGIS/SINBA\nCFDI · FHIR · RENAPO\nNotificaciones", CORAL)
    add_textbox(
        s,
        Inches(0.5),
        Inches(4.8),
        Inches(12.3),
        Inches(1.8),
        "Contrato: fallar una capacidad externa no tumba la consulta.\nSí: eventos de dominio + empuje al cliente · No: CQRS completo, consistencia eventual en lecturas clínicas, microservicios.",
        14,
        False,
        INK,
    )

    # ---- 12 Integrations ----
    s = new()
    banner(s, "Integraciones", "Fijo vs feature flags")
    card(
        s,
        Inches(0.45),
        Inches(1.35),
        Inches(6.1),
        Inches(5.35),
        "DGIS / SINBA — fijo (ADR-023)",
        "• Siempre en el producto (todo tenant / instalación)\n• Con red: genera + envía vía outbox\n• Sin red: genera + queda en outbox\n• Demo: misma tubería, destino no productivo\n• On-prem: módulo presente; cliente provee Internet para envío\n• No existe SKU “sin DGIS”",
        TEAL,
        13,
    )
    card(
        s,
        Inches(6.8),
        Inches(1.35),
        Inches(6.1),
        Inches(5.35),
        "Feature flags (solo estas)",
        "• CFDI — facturación electrónica\n• FHIR — interoperabilidad voluntaria (≠ NOM-024)\n• RENAPO — consulta de identidad\n\nSin Internet o flag off: no bloquean atención.\nConformidad NOM-024 = GIIS/DGIS, no FHIR.",
        AMBER,
        13,
    )

    # ---- 13 Multi-tenant / security ----
    s = new()
    banner(s, "Multi-tenant, persistencia y seguridad", "API = fuente de verdad de autorización")
    for i, (t, b) in enumerate(
        [
            ("TenantId", "En claims JWT y en todo SP de negocio. Sin BD-por-tenant."),
            ("SPs", "Única vía de persistencia de negocio: sp_{Entity}_{Action}."),
            ("Auth", "Refresh httpOnly (ADR-006 aprobado). MFA/dispositivos según fase."),
            ("Identidad", "Progresiva: sujeto ≠ episodio (ADR-016). CURP opcional."),
            ("Bloqueo", "Nada bloquea ingreso a urgencias (ADR-015)."),
            ("Modelo", "Debe poder decir “no sé” (doc 09: 87 hallazgos)."),
            ("PDF clínico", "Generado en servidor (ADR-009); no html2canvas como evidencia."),
            ("Firma", "Local + sello al sync (ADR-020); dictamen legal pendiente."),
            ("Rechazo diferido", "Bandeja por severidad; nunca borrar; corrección por addendum."),
        ]
    ):
        col, row = i % 3, i // 3
        card(s, Inches(0.4) + Inches(col * 4.25), Inches(1.3) + Inches(row * 1.85), Inches(4.1), Inches(1.7), t, b, body_size=12)

    # ---- 14 BI ----
    s = new()
    banner(s, "BI / dirección (ADR-013)", "Módulo en la app principal — no app aparte")
    for i, (t, b) in enumerate(
        [
            ("Ubicación", "features/bi dentro del frontend clínico; lazy load para no pesar urgencias."),
            ("Permisos", "Solo lectura; perfiles de dirección; API GET de agregados."),
            ("Anti-reID", "Solo agregados; supresión de celdas pequeñas; sin abrir expediente desde KPI."),
            ("Offline", "Sin cola/caché clínica del tablero. Auditoría de todo acceso."),
        ]
    ):
        col, row = i % 2, i // 2
        card(s, Inches(0.5) + Inches(col * 6.35), Inches(1.4) + Inches(row * 2.6), Inches(6.1), Inches(2.4), t, b, body_size=14)

    # ---- 15 Hosting ----
    s = new()
    banner(s, "Hospedaje y on-prem", "ADR-021 / ADR-022")
    for i, (t, lines, color) in enumerate(
        [
            (
                "Shared (demo/ventas)",
                "Datos sintéticos únicamente\nBanner demo\nFlags CFDI/FHIR/RENAPO off\nDGIS → destino no productivo\nSin deploy escalonado fiable",
                AMBER,
            ),
            (
                "Dedicado (producción)",
                "Obligatorio antes del 1er paciente real\nWorker durable\nPITR / DR\nDeploy escalonado de API\nResidencia según contrato",
                TEAL,
            ),
            (
                "On-prem cliente",
                "Mismo artefacto (API+Worker+SQL)\nN estaciones en LAN\nSin sync expediente a cloud\nMulti-sede: N installs o agregados\nPortabilidad desde Fase 0",
                NAVY,
            ),
        ]
    ):
        left = Inches(0.4) + Inches(i * 4.25)
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.4), Inches(4.1), Inches(5.2))
        fill_shape(shp, CARD_BG)
        shp.line.color.rgb = LINE
        hdr = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.4), Inches(4.1), Inches(0.7))
        fill_shape(hdr, color)
        add_textbox(s, left + Inches(0.15), Inches(1.55), Inches(3.8), Inches(0.45), t, 15, True, WHITE, PP_ALIGN.CENTER)
        add_textbox(s, left + Inches(0.25), Inches(2.4), Inches(3.6), Inches(3.8), lines, 14, False, INK)

    # ---- 16 ADRs ----
    s = new()
    banner(s, "ADRs vigentes (resumen)", "Estado 2026-08-23")
    adrs = [
        "001 Core+PWA sin Edge",
        "003 SPA=PWA",
        "005 ULID en borde",
        "006 Refresh httpOnly ✓",
        "007 Continuidad en dispositivo",
        "009 PDF en servidor",
        "012 API idéntica cloud/on-prem",
        "013 BI en app principal",
        "014 Escritura siempre local",
        "015 Nada bloquea ingreso",
        "016 Identidad progresiva",
        "017 Registro dispositivos",
        "018 Outbox + EDA acotada",
        "019 Cobro/farmacia offline",
        "020 Firma local + sello",
        "021 Portabilidad on-prem",
        "022 Demo vs producción",
        "023 DGIS/SINBA fijo",
    ]
    for i, a in enumerate(adrs):
        col, row = i % 3, i // 3
        left = Inches(0.45) + Inches(col * 4.25)
        top = Inches(1.3) + Inches(row * 0.9)
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(4.1), Inches(0.75))
        fill_shape(shp, CARD_BG)
        shp.line.color.rgb = LINE
        add_textbox(s, left + Inches(0.15), top + Inches(0.18), Inches(3.8), Inches(0.4), f"ADR-{a}", 13, False, INK)
    add_textbox(s, Inches(0.5), Inches(6.7), Inches(12), Inches(0.35), "Obsoletos: ADR-010 Edge Express · ADR-011 modalidad Edge por sucursal", 11, False, CORAL)

    # ---- 17 Non-promises ----
    s = new()
    banner(s, "Contrato de no-promesas", "Para evitar deuda de expectativas en arquitectura")
    nos = [
        "Multiusuario offline equivalente a servidor LAN en clínica",
        "24/7 sin degradación si cae la BD central",
        "Certificación NOM-024 automática",
        "CFDI/FHIR/RENAPO sin Internet o con flag off (DGIS: solo se difiere el envío)",
        "Estupefacientes / psicotrópicos (impedir hasta norma + decisión)",
        "Demo creíble sin backend real (sin mocks que tapen integración)",
        "Sync vivo de expediente on-prem ↔ cloud",
        "Redis / broker / microservicios en MVP",
    ]
    bullets(s, Inches(0.7), Inches(1.4), Inches(12), Inches(5.5), nos, 16)

    # ---- 18 Repo / phases ----
    s = new()
    banner(s, "Estructura objetivo y fases", "Implementación no arranca sin autorización Fase 0")
    code = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(1.35), Inches(5.8), Inches(5.3))
    fill_shape(code, CODE_BG)
    add_textbox(
        s,
        Inches(0.7),
        Inches(1.6),
        Inches(5.3),
        Inches(4.8),
        "MedicalCore/\n  backend/   Api, Worker, Business,\n             DataAccess, Models, Common,\n             database/\n  frontend/  PWA + features/bi\n  e2e/       Playwright vs stack real\n  tools/     seed-demo, diagnose\n  docs/      análisis + prototipo",
        14,
        False,
        RGBColor(0xC8, 0xE4, 0xE8),
        font="Consolas",
    )
    phases = [
        ("0", "Fundación real + cola + tenant + vertical E2E"),
        ("D", "Demo ventas: stack real + seed sintético"),
        ("1", "Núcleo clínico + live"),
        ("2", "Caja / CFDI / farmacia / estudios"),
        ("3", "Privacidad + módulo BI"),
        ("4", "DGIS productivo + FHIR opcional"),
        ("5", "Hospital / instalador on-prem producto → aparte"),
    ]
    for i, (n, t) in enumerate(phases):
        top = Inches(1.35) + Inches(i * 0.72)
        nb = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.5), top, Inches(0.7), Inches(0.55))
        fill_shape(nb, TEAL if n != "5" else CORAL)
        add_textbox(s, Inches(6.5), top + Inches(0.08), Inches(0.7), Inches(0.4), n, 14, True, WHITE, PP_ALIGN.CENTER)
        add_textbox(s, Inches(7.4), top + Inches(0.1), Inches(5.4), Inches(0.45), t, 13, False, INK)

    # ---- 19 Decision checklist ----
    s = new()
    banner(s, "Checklist para revisión de arquitectura", "Preguntas que un arquitecto debe poder responder “sí”")
    checks = [
        "¿Escritura siempre pasa por cola local + idempotencia?",
        "¿Lectura online viene del servidor (no de caché como verdad)?",
        "¿Ningún módulo reintroduce Edge o sync de expediente?",
        "¿Todo SP de negocio recibe TenantId?",
        "¿DGIS/SINBA existe aunque no haya red (outbox)?",
        "¿CFDI/FHIR/RENAPO están detrás de flag y no bloquean atención?",
        "¿BI es módulo de agregados sin PHI de paciente?",
        "¿On-prem = mismo artefacto, sin consolidar expediente al centro?",
        "¿Ambiente shared = solo sintético; dedicado antes de PHI?",
        "¿Estupefacientes impedidos hasta decisión explícita?",
    ]
    bullets(s, Inches(0.7), Inches(1.35), Inches(12), Inches(5.5), checks, 15)

    # ---- 20 Cierre ----
    s = new()
    fill_shape(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height), NAVY)
    fill_shape(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(5.7), prs.slide_width, Inches(1.8)), TEAL)
    add_textbox(s, Inches(0.8), Inches(1.8), Inches(11.5), Inches(0.6), "Resumen en una línea", 28, True, WHITE, font="Consolas")
    add_textbox(
        s,
        Inches(0.8),
        Inches(2.7),
        Inches(11.5),
        Inches(2.0),
        "Un Core multi-tenant con SPs; un cliente SPA/PWA con cola durable;\noutbox hacia el exterior; continuidad por dispositivo;\nsin Edge y con degradación explícita.",
        18,
        False,
        RGBColor(0xC8, 0xE4, 0xE6),
    )
    add_textbox(
        s,
        Inches(0.8),
        Inches(5.95),
        Inches(11.5),
        Inches(1.1),
        "Referencias: 12-propuesta-final.md · 03-arquitectura-propuesta.md · AGENTS.md\nPróximo artefacto de ingeniería: ADRs formales + diagramas C4 al arrancar Fase 0",
        13,
        False,
        WHITE,
    )

    total = len(slides)
    for idx, sl in enumerate(slides, start=1):
        if idx in (1, total):
            continue
        footer(sl, idx, total)

    prs.save(OUT)
    print(f"Generado: {OUT}")
    print(f"Diapositivas: {total}")


if __name__ == "__main__":
    build()
