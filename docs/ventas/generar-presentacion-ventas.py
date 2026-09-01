# -*- coding: utf-8 -*-
"""Genera la presentación comercial MediCore (lenguaje no técnico)."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("MediCore-Presentacion-Comercial.pptx")

NAVY = RGBColor(0x0B, 0x2C, 0x3D)
TEAL = RGBColor(0x0D, 0x7A, 0x7A)
TEAL_LIGHT = RGBColor(0xE6, 0xF4, 0xF4)
CORAL = RGBColor(0xC4, 0x5C, 0x26)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1A, 0x2A, 0x32)
MUTED = RGBColor(0x5A, 0x6B, 0x73)
LINE = RGBColor(0xD0, 0xDC, 0xE0)
CARD_BG = RGBColor(0xF7, 0xFA, 0xFB)


def set_run(run, text, size=18, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    set_run(r, text, size, bold, color, font)
    return box


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def banner(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
    fill_shape(bar, NAVY)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), Inches(13.333), Inches(0.08))
    fill_shape(accent, TEAL)
    add_textbox(slide, Inches(0.55), Inches(0.28), Inches(12), Inches(0.55), title, 28, True, WHITE)
    if subtitle:
        add_textbox(
            slide,
            Inches(0.55),
            Inches(0.72),
            Inches(12),
            Inches(0.35),
            subtitle,
            14,
            False,
            RGBColor(0xB8, 0xD4, 0xD8),
        )


def footer(slide, page, total):
    add_textbox(
        slide,
        Inches(0.55),
        Inches(7.15),
        Inches(10),
        Inches(0.3),
        "MediCore  ·  Confidencial — uso comercial",
        11,
        False,
        MUTED,
    )
    add_textbox(
        slide,
        Inches(11.2),
        Inches(7.15),
        Inches(1.5),
        Inches(0.3),
        f"{page} / {total}",
        11,
        False,
        MUTED,
        PP_ALIGN.RIGHT,
    )


def card(slide, left, top, width, height, title, body, title_color=TEAL):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill_shape(shp, CARD_BG)
    shp.line.color.rgb = LINE
    shp.adjustments[0] = 0.08
    add_textbox(slide, left + Inches(0.2), top + Inches(0.15), width - Inches(0.35), Inches(0.4), title, 15, True, title_color)
    add_textbox(slide, left + Inches(0.2), top + Inches(0.55), width - Inches(0.35), height - Inches(0.7), body, 13, False, INK)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slides = []

    def new():
        sl = prs.slides.add_slide(blank)
        slides.append(sl)
        return sl

    # 1 Portada
    s = new()
    fill_shape(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height), NAVY)
    fill_shape(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(5.55), prs.slide_width, Inches(1.95)), TEAL)
    add_textbox(s, Inches(0.85), Inches(1.7), Inches(11.5), Inches(0.95), "MediCore", 54, True, WHITE, font="Georgia")
    add_textbox(
        s,
        Inches(0.85),
        Inches(2.75),
        Inches(11.5),
        Inches(1.2),
        "El sistema que acompaña la atención médica\nde su clínica, todos los días.",
        24,
        False,
        RGBColor(0xC8, 0xE4, 0xE6),
    )
    add_textbox(
        s,
        Inches(0.85),
        Inches(5.85),
        Inches(11.5),
        Inches(1.0),
        "Presentación comercial  ·  Clínicas ambulatorias y servicios de urgencias\nMéxico",
        16,
        False,
        WHITE,
    )

    # 2 Agenda
    s = new()
    banner(s, "Agenda", "Qué veremos juntos")
    agenda = [
        ("01", "El reto de operar una clínica hoy"),
        ("02", "Qué es MediCore y para quién"),
        ("03", "Cómo ayuda en el día a día clínico"),
        ("04", "Caja, farmacia y operación"),
        ("05", "Cuando falla el Internet"),
        ("06", "Varias sucursales y su propia marca"),
        ("07", "Información para la dirección"),
        ("08", "Cumplimiento y formas de instalación"),
        ("09", "Siguiente paso"),
    ]
    for i, (num, txt) in enumerate(agenda):
        col, row = i % 2, i // 2
        left = Inches(0.7) + Inches(col * 6.2)
        top = Inches(1.55) + Inches(row * 0.95)
        nbox = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(0.7), Inches(0.55))
        fill_shape(nbox, TEAL)
        nbox.adjustments[0] = 0.2
        add_textbox(s, left, top + Inches(0.08), Inches(0.7), Inches(0.4), num, 16, True, WHITE, PP_ALIGN.CENTER)
        add_textbox(s, left + Inches(0.9), top + Inches(0.12), Inches(4.9), Inches(0.4), txt, 16, False, INK)

    # 3 Reto
    s = new()
    banner(s, "El reto de muchas clínicas", "Lo que suele pasar en la operación diaria")
    for i, (t, b) in enumerate(
        [
            (
                "Papel y sistemas sueltos",
                "La información del paciente vive en carpetas, hojas de cálculo o programas que no se comunican.",
            ),
            (
                "Urgencias sin tiempo",
                "Llega alguien sin documentos o inconsciente. El software no debería frenar la atención.",
            ),
            (
                "Internet inestable",
                "Si se cae la red, la clínica no puede dejar de atender, cobrar ni surtir lo esencial.",
            ),
            (
                "Varias sucursales",
                "La dirección necesita ver el conjunto sin perder el control de cada sede.",
            ),
        ]
    ):
        col, row = i % 2, i // 2
        card(s, Inches(0.55) + Inches(col * 6.3), Inches(1.55) + Inches(row * 2.5), Inches(6.0), Inches(2.25), t, b)

    # 4 Qué es
    s = new()
    banner(s, "Qué es MediCore", "En una frase")
    q = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.65), Inches(11.9), Inches(2.05))
    fill_shape(q, TEAL_LIGHT)
    q.line.color.rgb = TEAL
    q.adjustments[0] = 0.06
    add_textbox(
        s,
        Inches(1.0),
        Inches(2.0),
        Inches(11.3),
        Inches(1.4),
        "MediCore es un sistema web para clínicas ambulatorias\ny servicios de urgencias: unifica la atención, el expediente,\nla caja y la farmacia en una sola herramienta.",
        21,
        True,
        NAVY,
        PP_ALIGN.CENTER,
    )
    for i, (t, b) in enumerate(
        [
            (
                "Una sola plataforma",
                "Recepción, triage, consulta, urgencias, recetas, cobro y farmacia en el mismo flujo.",
            ),
            (
                "Pensado para México",
                "Reportes de salud, facturación y cuidados de privacidad acordes a su operación.",
            ),
            ("Con su marca", "Logo, colores y nombre de su clínica o de su red de sucursales."),
        ]
    ):
        card(s, Inches(0.55) + Inches(i * 4.15), Inches(4.05), Inches(3.95), Inches(2.25), t, b)

    # 5 Para quién
    s = new()
    banner(s, "Para quién es", "Alcance claro, sin sorpresas")
    left_card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(1.5), Inches(6.0), Inches(5.1))
    fill_shape(left_card, CARD_BG)
    left_card.line.color.rgb = LINE
    left_card.adjustments[0] = 0.05
    add_textbox(s, Inches(0.85), Inches(1.7), Inches(5.4), Inches(0.45), "Ideal para", 20, True, TEAL)
    box = s.shapes.add_textbox(Inches(0.85), Inches(2.3), Inches(5.4), Inches(4.0))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(
        [
            "Clínicas privadas ambulatorias",
            "Redes con varias sucursales",
            "Consultorios con servicio de urgencias",
            "Operaciones que quieren expediente, caja y farmacia juntos",
            "Equipos que trabajan en computadora o tablet dentro de la clínica",
        ]
    ):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(8)
        r = p.add_run()
        set_run(r, f"•  {line}", 16, False, INK)

    right_card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(6.0), Inches(5.1))
    fill_shape(right_card, RGBColor(0xFD, 0xF4, 0xEE))
    right_card.line.color.rgb = RGBColor(0xE8, 0xC4, 0xA8)
    right_card.adjustments[0] = 0.05
    add_textbox(s, Inches(7.1), Inches(1.7), Inches(5.4), Inches(0.45), "Fuera de este alcance", 20, True, CORAL)
    box = s.shapes.add_textbox(Inches(7.1), Inches(2.3), Inches(5.4), Inches(4.0))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(
        [
            "Hospitalización completa, quirófanos o UCI (proyecto aparte)",
            "Portal del paciente como producto principal",
            "Medicamentos de alto control (estupefacientes): pendiente de definición normativa",
        ]
    ):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(8)
        r = p.add_run()
        set_run(r, f"•  {line}", 15, False, INK)

    # 6 Promesa
    s = new()
    banner(s, "Nuestra promesa", "Lo más importante del producto")
    big = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(2.0), Inches(10.9), Inches(2.4))
    fill_shape(big, NAVY)
    big.adjustments[0] = 0.06
    add_textbox(
        s,
        Inches(1.6),
        Inches(2.45),
        Inches(10.1),
        Inches(1.6),
        "La atención no se detiene.\nLo administrativo se completa después,\ncon rastro claro de lo ocurrido.",
        26,
        True,
        WHITE,
        PP_ALIGN.CENTER,
    )
    add_textbox(
        s,
        Inches(1.2),
        Inches(4.8),
        Inches(10.9),
        Inches(1.5),
        "En urgencias puede registrar y atender aunque falten datos;\nluego se completa la identidad y el cobro sin perder lo clínico.\nSiempre queda registrado quién hizo qué y cuándo.",
        16,
        False,
        MUTED,
        PP_ALIGN.CENTER,
    )

    # 7 Recorrido
    s = new()
    banner(s, "El recorrido del paciente", "De la llegada al egreso, en un solo sistema")
    steps = [
        ("1", "Recepción", "Alta o búsqueda del paciente"),
        ("2", "Triage", "Prioridad según gravedad"),
        ("3", "Atención", "Consulta o urgencias"),
        ("4", "Órdenes", "Receta, estudios, insumos"),
        ("5", "Caja", "Cobro y comprobante"),
        ("6", "Farmacia", "Surtido y control"),
    ]
    for i, (n, t, b) in enumerate(steps):
        left = Inches(0.45) + Inches(i * 2.15)
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.55), Inches(2.0), Inches(0.7), Inches(0.7))
        fill_shape(circ, TEAL)
        add_textbox(s, left + Inches(0.55), Inches(2.12), Inches(0.7), Inches(0.5), n, 20, True, WHITE, PP_ALIGN.CENTER)
        add_textbox(s, left, Inches(2.9), Inches(2.0), Inches(0.4), t, 15, True, NAVY, PP_ALIGN.CENTER)
        add_textbox(s, left, Inches(3.35), Inches(2.0), Inches(0.9), b, 12, False, MUTED, PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(1.85), Inches(2.3), Inches(0.55), Inches(0.08))
            fill_shape(line, LINE)
    for i, (t, b) in enumerate(
        [
            ("Todo queda en el expediente", "Notas, signos, diagnósticos y decisiones clínicas en un solo lugar."),
            ("Menos retrabajo", "Lo capturado en recepción y triage se reutiliza en consulta y caja."),
            ("Visibilidad en sala", "El equipo ve turnos y estados actualizados mientras hay conexión."),
        ]
    ):
        card(s, Inches(0.55) + Inches(i * 4.15), Inches(4.5), Inches(3.95), Inches(2.0), t, b)

    # 8 Clínico
    s = new()
    banner(s, "Módulos clínicos", "Lo que usa el personal de salud")
    mods = [
        (
            "Identificación flexible",
            "Puede iniciar con pocos datos y completar CURP u otros datos después. Ideal para urgencias.",
        ),
        (
            "Paciente no identificado",
            "Atiende con una etiqueta temporal clara (por ejemplo, señas del accidente) y luego une el expediente cuando se sepa quién es.",
        ),
        (
            "Triage",
            "Clasifica prioridad de atención. El color clínico no se confunde con los colores de su marca.",
        ),
        ("Consulta ambulatoria", "Nota clínica, antecedentes, signos vitales, plan e interconsulta."),
        ("Urgencias", "Ingreso rápido, seguimiento en sala y registro sin bloquear por papelería."),
        (
            "Receta e expediente",
            "Prescripción con datos del médico; nota firmada que no se altera una vez cerrada.",
        ),
    ]
    for i, (t, b) in enumerate(mods):
        col, row = i % 3, i // 3
        card(s, Inches(0.45) + Inches(col * 4.25), Inches(1.5) + Inches(row * 2.55), Inches(4.05), Inches(2.35), t, b)

    # 9 Operación
    s = new()
    banner(s, "Caja, farmacia y estudios", "La operación que acompaña la clínica")
    for i, (t, b) in enumerate(
        [
            (
                "Caja",
                "Cobros ligados a la atención. Si no hay Internet, puede emitir un recibo provisional y facturar después cuando vuelva la conexión.",
            ),
            (
                "Factura electrónica",
                "Cuando su clínica lo active: generación de CFDI acorde a su operación fiscal.",
            ),
            (
                "Farmacia",
                "Surtido de recetas e inventario. Si falta conexión, puede seguir surtiendo y reconciliar existencias después.",
            ),
            ("Estudios", "Órdenes de laboratorio o gabinete dentro del mismo flujo del paciente."),
            (
                "Rechazos y correcciones",
                "Si algo no pudo validarse al sincronizar, no se pierde: queda en una bandeja para revisar y corregir.",
            ),
            (
                "Usuarios y permisos",
                "Cada persona ve solo lo que le corresponde: recepción, médico, caja, farmacia o dirección.",
            ),
        ]
    ):
        col, row = i % 3, i // 3
        card(s, Inches(0.45) + Inches(col * 4.25), Inches(1.5) + Inches(row * 2.55), Inches(4.05), Inches(2.35), t, b)

    # 10 Offline
    s = new()
    banner(s, "Cuando falla el Internet", "La clínica sigue atendiendo")
    add_textbox(
        s,
        Inches(0.7),
        Inches(1.5),
        Inches(12),
        Inches(0.6),
        "MediCore está pensado para que un corte de red no detenga la atención en esa estación de trabajo.",
        17,
        False,
        INK,
    )
    for i, (t, b) in enumerate(
        [
            (
                "Con Internet",
                "Funciona completo: consulta en vivo, actualizaciones en sala y envío de reportes o facturas.",
            ),
            (
                "Sin Internet",
                "Esa computadora o tablet puede seguir capturando atención, cobro y surtido según reglas acordadas. Al volver la red, se envía lo pendiente.",
            ),
            (
                "Importante",
                "Cada estación trabaja con su propia cola. No es un “servidor de piso” compartido entre todas las PCs sin conexión.",
            ),
        ]
    ):
        card(
            s,
            Inches(0.55) + Inches(i * 4.15),
            Inches(2.3),
            Inches(3.95),
            Inches(3.6),
            t,
            b,
            TEAL if i < 2 else CORAL,
        )

    # 11 Multi-sucursal
    s = new()
    banner(s, "Varias sucursales, una misma casa", "Multi-clínica y su identidad visual")
    for i, (t, b) in enumerate(
        [
            (
                "Varias sedes",
                "Una misma organización puede operar varias sucursales con usuarios, horarios y operación propios.",
            ),
            (
                "Su marca (white-label)",
                "Logo, colores y nombre de su clínica. La identidad clínica (por ejemplo, colores de triage) se protege para no confundir al personal.",
            ),
            (
                "Instalación en la nube",
                "Sus equipos se conectan por navegador. Ideal si no quiere administrar servidores.",
            ),
            (
                "En su propio servidor",
                "Si prefiere infraestructura local: mismo sistema instalado en su servidor, con varios clientes en la red de la clínica.",
            ),
        ]
    ):
        col, row = i % 2, i // 2
        card(s, Inches(0.55) + Inches(col * 6.3), Inches(1.5) + Inches(row * 2.55), Inches(6.0), Inches(2.35), t, b)

    # 12 BI
    s = new()
    banner(s, "Información para la dirección", "Ver el negocio sin exponer el expediente")
    q = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.6), Inches(11.9), Inches(1.5))
    fill_shape(q, TEAL_LIGHT)
    q.line.color.rgb = TEAL
    add_textbox(
        s,
        Inches(1.0),
        Inches(1.9),
        Inches(11.3),
        Inches(1.0),
        "Un módulo de indicadores dentro de la misma aplicación — no es un sistema aparte.\nSolo números agregados, con permisos de lectura.",
        18,
        True,
        NAVY,
        PP_ALIGN.CENTER,
    )
    for i, (t, b) in enumerate(
        [
            (
                "Qué sí ve",
                "Volúmenes de atención, tiempos, ocupación, ingresos agregados y tendencias por sucursal.",
            ),
            ("Qué no hace", "No abre el expediente clínico desde un indicador. Protege la privacidad del paciente."),
            ("Para quién", "Dirección, administración y perfiles autorizados. Cada acceso queda registrado."),
        ]
    ):
        card(s, Inches(0.55) + Inches(i * 4.15), Inches(3.5), Inches(3.95), Inches(2.7), t, b)

    # 13 Cumplimiento
    s = new()
    banner(s, "Cumplimiento y reportes", "En lenguaje sencillo")
    for i, (t, b) in enumerate(
        [
            (
                "Reportes de salud (DGIS / SINBA)",
                "Incluidos en el producto. Si no hay Internet, se guardan y se envían después. No hay que “activar un módulo extra” para tenerlos.",
            ),
            (
                "Privacidad de datos",
                "Diseñado para cuidar información sensible de pacientes, con control de accesos y registro de quién consultó qué.",
            ),
            (
                "Integraciones opcionales",
                "Según su contrato: facturación electrónica, consulta de identidad oficial u otras conexiones. Se activan cuando las necesite.",
            ),
            (
                "Documentación clínica",
                "Notas y registros pensados para el trabajo diario de la clínica, con sello de tiempo y autoría.",
            ),
        ]
    ):
        col, row = i % 2, i // 2
        card(s, Inches(0.55) + Inches(col * 6.3), Inches(1.5) + Inches(row * 2.55), Inches(6.0), Inches(2.35), t, b)

    # 14 Cómo se usa
    s = new()
    banner(s, "Cómo lo usa su equipo", "Simple de adoptar")
    for i, (t, b) in enumerate(
        [
            (
                "Desde el navegador",
                "Entra con su usuario desde la computadora de la clínica. No depende de un programa distinto por área.",
            ),
            (
                "También como app",
                "Puede instalarse en el dispositivo para una experiencia más ágil (misma herramienta).",
            ),
            (
                "Dispositivos de confianza",
                "Las estaciones que guardan información clínica se registran: más control y menos sorpresas.",
            ),
            (
                "Capacitación por rol",
                "Recepción, enfermería, médicos, caja y farmacia ven pantallas acordes a su trabajo.",
            ),
        ]
    ):
        col, row = i % 2, i // 2
        card(s, Inches(0.55) + Inches(col * 6.3), Inches(1.5) + Inches(row * 2.55), Inches(6.0), Inches(2.35), t, b)

    # 15 Beneficios
    s = new()
    banner(s, "Beneficios que sentirá su clínica", "Por qué conviene MediCore")
    benefits = [
        ("Menos fricción en urgencias", "Atiende primero; completa papeles después."),
        ("Un solo lugar para todo", "Menos sistemas, menos capturas repetidas."),
        ("Continuidad ante cortes", "La estación sigue trabajando y luego se sincroniza."),
        ("Visión de dirección", "Indicadores sin meterse al expediente."),
        ("Su imagen, su red", "Marca propia y varias sucursales."),
        ("Listo para reportar", "Capacidad de reportes de salud desde el producto."),
    ]
    for i, (t, b) in enumerate(benefits):
        col, row = i % 3, i // 3
        left = Inches(0.45) + Inches(col * 4.25)
        top = Inches(1.55) + Inches(row * 2.55)
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(4.05), Inches(2.3))
        fill_shape(shp, CARD_BG)
        shp.line.color.rgb = LINE
        shp.adjustments[0] = 0.08
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.12), Inches(2.3))
        fill_shape(bar, TEAL)
        add_textbox(s, left + Inches(0.35), top + Inches(0.45), Inches(3.5), Inches(0.6), t, 16, True, NAVY)
        add_textbox(s, left + Inches(0.35), top + Inches(1.15), Inches(3.5), Inches(0.8), b, 14, False, MUTED)

    # 16 Modalidades
    s = new()
    banner(s, "Modalidades de operación", "Elija cómo quiere trabajar")
    modalities = [
        (
            "En la nube",
            [
                "Nosotros (o su proveedor) hospedan el sistema",
                "Usted se conecta por Internet",
                "Ideal para empezar rápido",
                "Varias sucursales en un mismo entorno",
            ],
            TEAL,
        ),
        (
            "En su servidor",
            [
                "El sistema corre en infraestructura suya",
                "Varios equipos en su red local",
                "Menos dependencia del Internet externo",
                "Misma experiencia de usuario",
            ],
            TEAL,
        ),
        (
            "Según su madurez",
            [
                "Demo con datos de ejemplo",
                "Puesta en marcha con acompañamiento",
                "Activación de facturación u otras conexiones cuando las necesite",
                "Evolución a hospital: proyecto aparte",
            ],
            NAVY,
        ),
    ]
    for i, (t, lines, color) in enumerate(modalities):
        left = Inches(0.45) + Inches(i * 4.25)
        top = Inches(1.5)
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(4.05), Inches(5.1))
        fill_shape(shp, CARD_BG)
        shp.line.color.rgb = LINE
        shp.adjustments[0] = 0.05
        hdr = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(4.05), Inches(0.7))
        fill_shape(hdr, color)
        add_textbox(s, left + Inches(0.2), top + Inches(0.18), Inches(3.65), Inches(0.45), t, 18, True, WHITE, PP_ALIGN.CENTER)
        box = s.shapes.add_textbox(left + Inches(0.3), Inches(2.5), Inches(3.5), Inches(3.8))
        tf = box.text_frame
        tf.word_wrap = True
        for j, line in enumerate(lines):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_before = Pt(10)
            r = p.add_run()
            set_run(r, f"•  {line}", 14, False, INK)

    # 17 Preguntas
    s = new()
    banner(s, "Para armar una propuesta a su medida", "Preguntas que nos ayudan a cotizar bien")
    qs = [
        "¿Cuántas sucursales y cuántos usuarios por rol?",
        "¿Operan urgencias las 24 horas o en horarios definidos?",
        "¿Necesitan facturación electrónica desde el día uno?",
        "¿Prefieren nube o servidor en sus instalaciones?",
        "¿Tienen farmacia propia e inventario?",
        "¿Qué sistemas usan hoy (si aplica) y qué duele más?",
    ]
    box = s.shapes.add_textbox(Inches(0.9), Inches(1.6), Inches(11.5), Inches(5.0))
    tf = box.text_frame
    tf.word_wrap = True
    for i, qtext in enumerate(qs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(12)
        r = p.add_run()
        set_run(r, f"{i + 1}.  {qtext}", 18, False, INK)

    # 18 Cierre
    s = new()
    fill_shape(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height), NAVY)
    fill_shape(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(5.7), prs.slide_width, Inches(1.8)), TEAL)
    add_textbox(s, Inches(0.85), Inches(1.8), Inches(11.5), Inches(0.7), "Siguiente paso", 36, True, WHITE, font="Georgia")
    add_textbox(
        s,
        Inches(0.85),
        Inches(2.7),
        Inches(11.5),
        Inches(2.0),
        "1.  Recorrido guiado del sistema (demo)\n2.  Entrevista de operación de su clínica\n3.  Propuesta de alcance, tiempos y modalidad\n4.  Plan de arranque por etapas",
        20,
        False,
        RGBColor(0xC8, 0xE4, 0xE6),
    )
    add_textbox(
        s,
        Inches(0.85),
        Inches(6.0),
        Inches(11.5),
        Inches(1.0),
        "MediCore  ·  Atención que no se detiene\n¿Conversamos sobre su clínica?",
        18,
        True,
        WHITE,
    )

    total = len(slides)
    for idx, sl in enumerate(slides, start=1):
        if idx in (1, total):
            continue
        footer(sl, idx, total)

    prs.save(OUT)
    print(f"Generado: {OUT}")
    print(f"Diapositivas: {total}")


if __name__ == "__main__":
    build()
